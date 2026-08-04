"""
Citation Tools -- one app, two modes you can flip between:

  * Renumber  -- numbered citations (superscript / [n] / (n)); reorder, insert
                 or delete references and renumber everything.
  * Convert   -- turn author-date citations (Smith, 2020) into numbered
                 superscripts, with a review-before-you-commit step.

Import a document once; flipping modes re-reads the same file for that mode.
Runs locally or on a shared host. No secrets; the original file is never
modified.
"""

import os
import re
import threading
import uuid
import webbrowser

import docx
from flask import Flask, jsonify, render_template, request, send_file

import authordate as ad
import citations
import mendeley

HERE = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(HERE, "uploads")
EXPORT_DIR = os.path.join(HERE, "exports")
SAMPLE_NUM = os.path.join(HERE, "samples", "sample-article.docx")
SAMPLE_AD = os.path.join(HERE, "samples", "sample-authordate.docx")
PORT = int(os.environ.get("PORT", "5057"))

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(EXPORT_DIR, exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 40 * 1024 * 1024

_TOKEN_RE = re.compile(r"^[a-f0-9]{32}$")


def _resolve(token, default_sample):
    if not token or token == "sample":
        return default_sample if os.path.exists(default_sample) else None
    if _TOKEN_RE.match(token):
        p = os.path.join(UPLOAD_DIR, token + ".docx")
        return p if os.path.exists(p) else None
    return None


def _safe_base(filename):
    name = os.path.basename(filename or "document.docx")
    name = re.sub(r"\s*\(demo\)\s*", "", name)
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name)
    return os.path.splitext(name)[0] or "document"


def _safe_style(v):
    return v if v in citations.STYLES else None


# --------------------------------------------------------------------------
# Serialisers
# --------------------------------------------------------------------------

def num_dict(a, token):
    return {
        "token": token, "filename": a.filename, "blocks": a.blocks,
        "occurrences": a.occurrences, "references": a.references, "issues": a.issues,
        "style": a.style, "style_counts": a.style_counts,
        "counts": {
            "references": len(a.references), "citations": len(a.occurrences),
            "body": sum(1 for o in a.occurrences if o["location"] == "body"),
            "table": sum(1 for o in a.occurrences if o["location"] == "table"),
            "footnote": sum(1 for o in a.occurrences if o["location"] == "footnote"),
        },
    }


def ad_dict(document, filename, token):
    refs = ad.parse_references(list(document.paragraphs))
    occ, blocks = ad.scan(document)
    tokens = ad.propose(refs, occ)
    duplicates = citations.find_duplicate_groups([(r["id"], r["text"]) for r in refs])
    return {
        "token": token, "filename": filename, "blocks": blocks,
        "references": [{"id": r["id"], "surname": r["surname"],
                        "years": sorted(r["years"]), "text": r["text"]} for r in refs],
        "tokens": tokens,
        "duplicates": duplicates,
        "summary": {
            "citations": sum(t["count"] for t in tokens), "distinct": len(tokens),
            "matched": sum(1 for t in tokens if t["status"] == "matched"),
            "review": sum(1 for t in tokens if t["status"] == "review"),
            "unmatched": sum(1 for t in tokens if t["status"] == "unmatched"),
            "references": len(refs),
            "duplicates": len(duplicates),
        },
    }


# --------------------------------------------------------------------------
# Routes
# --------------------------------------------------------------------------

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/upload", methods=["POST"])
def upload():
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error": "No file received."}), 400
    if not f.filename.lower().endswith(".docx"):
        return jsonify({"error": "Please choose a .docx file."}), 400
    token = uuid.uuid4().hex
    dest = os.path.join(UPLOAD_DIR, token + ".docx")
    f.save(dest)
    # Mendeley-managed docs: flatten citations to numbered so they're readable.
    if mendeley.is_mendeley(dest):
        try:
            rep = mendeley.flatten(dest, dest)
            return jsonify({"token": token, "filename": f.filename,
                            "mendeley": rep})
        except Exception as exc:
            return jsonify({"token": token, "filename": f.filename,
                            "mendeley_error": str(exc)})
    return jsonify({"token": token, "filename": f.filename})


# ---- Renumber mode ----

@app.route("/api/num/analysis")
def num_analysis():
    token = request.args.get("token") or "sample"
    src = _resolve(token, SAMPLE_NUM)
    if src is None:
        return jsonify({"error": "Session expired -- please re-import."}), 400
    name = request.args.get("filename") or ("sample-article.docx (demo)"
                                            if token == "sample" else "document.docx")
    try:
        a = citations.analyze(src, name, style=_safe_style(request.args.get("style")))
    except Exception as exc:
        return jsonify({"error": "Could not read that document: %s" % exc}), 400
    return jsonify(num_dict(a, token))


@app.route("/api/num/renumber", methods=["POST"])
def num_renumber():
    p = request.get_json(force=True, silent=True) or {}
    src = _resolve(p.get("token"), SAMPLE_NUM)
    if src is None:
        return jsonify({"error": "Session expired -- please re-import."}), 400
    if not p.get("order"):
        return jsonify({"error": "No reference order supplied."}), 400
    out_name = _safe_base(p.get("filename")) + "__renumbered.docx"
    try:
        rep = citations.apply_renumber(src, p["order"], os.path.join(EXPORT_DIR, out_name),
                                       style=_safe_style(p.get("style")))
    except Exception as exc:
        return jsonify({"error": "Renumber failed: %s" % exc}), 400
    return jsonify({"ok": True, "download": "/download/" + out_name,
                    "change_log": rep["change_log"], "inserted": rep["inserted"],
                    "deleted": rep["deleted"], "warnings": rep["warnings"]})


# ---- Convert mode ----

@app.route("/api/ad/analysis")
def ad_analysis():
    token = request.args.get("token") or "sample"
    src = _resolve(token, SAMPLE_AD)
    if src is None:
        return jsonify({"error": "Session expired -- please re-import."}), 400
    name = request.args.get("filename") or ("sample-authordate.docx (demo)"
                                            if token == "sample" else "document.docx")
    try:
        return jsonify(ad_dict(docx.Document(src), name, token))
    except Exception as exc:
        return jsonify({"error": "Could not read that document: %s" % exc}), 400


@app.route("/api/ad/convert", methods=["POST"])
def ad_convert():
    p = request.get_json(force=True, silent=True) or {}
    src = _resolve(p.get("token"), SAMPLE_AD)
    if src is None:
        return jsonify({"error": "Session expired -- please re-import."}), 400
    mapping = {k: (int(v) if (v is not None and v != "") else None)
               for k, v in (p.get("mapping") or {}).items()}
    document = docx.Document(src)
    refs = ad.parse_references(list(document.paragraphs))
    occ = ad.detect_citations(document)
    try:
        rep = ad.apply_with_mapping(document, refs, occ, mapping)
    except Exception as exc:
        return jsonify({"error": "Conversion failed: %s" % exc}), 400
    out_name = _safe_base(p.get("filename")) + "__numbered.docx"
    document.save(os.path.join(EXPORT_DIR, out_name))
    return jsonify({"ok": True, "download": "/download/" + out_name,
                    "converted": rep["converted"], "left": rep["left"],
                    "numbered_refs": rep["numbered_refs"], "warnings": rep["warnings"]})


@app.route("/download/<path:name>")
def download(name):
    safe = os.path.basename(name)
    path = os.path.join(EXPORT_DIR, safe)
    if not os.path.exists(path):
        return "Not found", 404
    return send_file(path, as_attachment=True, download_name=safe)


def _open():
    webbrowser.open("http://127.0.0.1:%d/" % PORT)


if __name__ == "__main__":
    print("\n  Citation Tools running at  http://127.0.0.1:%d/\n" % PORT)
    threading.Timer(1.0, _open).start()
    app.run(host="127.0.0.1", port=PORT, debug=False)
